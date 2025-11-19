import os
import uuid
import shutil
import markdown
import json
import subprocess
from pathlib import Path
from typing import List, Optional
from datetime import datetime
from fastapi import FastAPI, UploadFile, File, Request, Form, Query, BackgroundTasks
from fastapi.responses import JSONResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from jinja2 import Environment, FileSystemLoader
from PIL import Image
from io import BytesIO
import traceback

from . import analyzer

BASE_DIR = Path(__file__).resolve().parents[1]
UPLOAD_DIR = BASE_DIR / "uploads"
SITE_DIR = BASE_DIR / "site"
IMAGES_DIR = SITE_DIR / "images"
# templates are located next to this module: app/templates
TEMPLATES_DIR = Path(__file__).resolve().parent / "templates"

os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(SITE_DIR, exist_ok=True)
os.makedirs(IMAGES_DIR, exist_ok=True)
# don't create templates dir at runtime; templates are part of the source tree

app = FastAPI(title="MPNN Blog Scaffold")
env = Environment(loader=FileSystemLoader(str(TEMPLATES_DIR)))

app.mount("/site", StaticFiles(directory=str(SITE_DIR)), name="site")


def save_upload_file(upload_file: UploadFile, dest: Path) -> Path:
    with open(dest, "wb") as f:
        shutil.copyfileobj(upload_file.file, f)
    return dest


def pptx_to_images(pptx_path: Path, slug: str) -> List[str]:
    """Convert each slide of a PPTX to a PNG image using LibreOffice."""
    img_paths = []
    output_dir = IMAGES_DIR / slug
    os.makedirs(output_dir, exist_ok=True)
    
    # Use a temporary directory for conversion artifacts
    temp_dir = output_dir / "temp"
    os.makedirs(temp_dir, exist_ok=True)
    
    pdf_path_stem = temp_dir / pptx_path.stem
    
    try:
        # 1. Convert PPTX to PDF using LibreOffice
        cmd_convert = [
            "soffice",
            "--headless",
            "--convert-to", "pdf:writer_pdf_Export",
            "--outdir",
            str(temp_dir),
            str(pptx_path)
        ]
        print(f"[INFO] Converting PPTX to PDF: {' '.join(cmd_convert)}")
        subprocess.run(cmd_convert, check=True, capture_output=True, timeout=180)
        
        # Find the generated PDF, as soffice might change the name
        generated_pdfs = list(temp_dir.glob("*.pdf"))
        if not generated_pdfs:
            raise FileNotFoundError("LibreOffice did not produce a PDF file from the PPTX.")
        pdf_path = generated_pdfs[0]

        # 2. Convert the resulting PDF to images
        cmd_images = [
            "pdftoppm",
            "-png",
            "-r", "150", # Set resolution to 150 DPI
            str(pdf_path),
            str(output_dir / "slide") # produces slide-1.png, slide-2.png etc
        ]
        print(f"[INFO] Converting PDF to images: {' '.join(cmd_images)}")
        subprocess.run(cmd_images, check=True, capture_output=True, timeout=180)
        
        # Collect generated image paths
        for f in sorted(output_dir.glob("*.png")):
            img_paths.append(f"/site/images/{slug}/{f.name}")
            
    except (subprocess.CalledProcessError, FileNotFoundError, subprocess.TimeoutExpired) as e:
        print(f"[ERROR] Failed to convert PPTX {pptx_path.name} via LibreOffice: {e}")
        if isinstance(e, subprocess.CalledProcessError):
            print(f"[ERROR] Stderr: {e.stderr.decode(errors='ignore')}")
            print(f"[ERROR] Stdout: {e.stdout.decode(errors='ignore')}")
        # Fallback to placeholder
        try:
            from pptx import Presentation
            prs = Presentation(pptx_path)
            for i in range(len(prs.slides)):
                placeholder = Image.new('RGB', (800, 600), color = (240, 240, 240))
                from PIL import ImageDraw
                d = ImageDraw.Draw(placeholder)
                d.text((10,10), f"Slide {i+1} (Conversion Failed)", fill=(0,0,0))
                img_filename = f"slide_{i+1}_fallback.png"
                img_path = output_dir / img_filename
                placeholder.save(img_path, "PNG")
                img_paths.append(f"/site/images/{slug}/{img_filename}")
        except Exception as fallback_e:
            print(f"[ERROR] PPTX fallback also failed: {fallback_e}")
    finally:
        # Clean up temporary conversion files
        shutil.rmtree(temp_dir, ignore_errors=True)
    
    return img_paths


def pdf_to_images(pdf_path: Path, slug: str) -> List[str]:
    """Convert each page of a PDF to a PNG image."""
    img_paths = []
    output_dir = IMAGES_DIR / slug
    os.makedirs(output_dir, exist_ok=True)
    
    try:
        # Using pdftoppm command line tool
        cmd = [
            "pdftoppm",
            "-png",
            str(pdf_path),
            str(output_dir / "page")
        ]
        subprocess.run(cmd, check=True, capture_output=True)
        
        # Collect generated image paths
        for f in sorted(output_dir.glob("*.png")):
            img_paths.append(f"/site/images/{slug}/{f.name}")
            
    except (subprocess.CalledProcessError, FileNotFoundError) as e:
        print(f"[ERROR] Failed to convert PDF {pdf_path.name} with pdftoppm: {e}")
        # Fallback to pdfplumber's visual debugging if available
        try:
            import pdfplumber
            with pdfplumber.open(pdf_path) as pdf:
                for i, page in enumerate(pdf.pages):
                    img = page.to_image(resolution=150)
                    img_filename = f"page_{i+1}.png"
                    img_path = output_dir / img_filename
                    img.save(img_path)
                    img_paths.append(f"/site/images/{slug}/{img_filename}")
        except Exception as e_fb:
            print(f"[ERROR] pdfplumber fallback for {pdf_path.name} also failed: {e_fb}")

    return img_paths


def parse_file(path: Path) -> str:
    name = path.name.lower()
    if name.endswith(".pdf"):
        try:
            import pdfplumber
            text_pages = []
            with pdfplumber.open(path) as pdf:
                for p in pdf.pages:
                    txt = p.extract_text(x_tolerance=2, y_tolerance=5)
                    if txt:
                        text_pages.append(txt)
            return "\n\n".join(text_pages)
        except Exception:
            return ""
    elif name.endswith(".docx"):
        try:
            import docx
            doc = docx.Document(path)
            return "\n\n".join(p.text for p in doc.paragraphs if p.text)
        except Exception:
            return ""
    elif name.endswith(".pptx"):
        try:
            from pptx import Presentation
            prs = Presentation(path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(shape.text)
            return "\n\n".join(texts)
        except Exception:
            return ""
    else:
        try:
            return path.read_text(encoding="utf-8")
        except Exception:
            return ""


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    # List generated posts with metadata (category, date, title, summary, tags)
    posts = []
    categories_data = {}  # {category: [posts]}
    for f in sorted(SITE_DIR.glob("*.json"), reverse=True):
        try:
            with open(f, "r", encoding="utf-8") as mf:
                data = json.load(mf)
            cat = data.get("category", "未分类")
            post_info = {
                "name": f.stem,
                "title": data.get("title", f.stem),
                "path": f"/site/{f.stem}.html",
                "category": cat,
                "date": data.get("date", ""),
                "tags": data.get("tags", []),
                "summary": data.get("summary", "")
            }
            posts.append(post_info)
            if cat not in categories_data:
                categories_data[cat] = []
            categories_data[cat].append(post_info)
        except Exception:
            # Fallback if JSON missing
            html_file = SITE_DIR / f"{f.stem}.html"
            if html_file.exists():
                posts.append({"name": f.stem, "title": f.stem, "path": f"/site/{f.stem}.html", "category": "未分类", "date": "", "tags": [], "summary": ""})
    tpl = env.get_template("index.html")
    return tpl.render(posts=posts, categories_data=categories_data)


@app.post("/upload")
async def upload(
    request: Request,
    files: List[UploadFile] = File(...),
    category: str = Form("未分类"),
    ts_override: Optional[str] = Form(None, alias="timestamp"),
    ts_q: Optional[str] = Query(None, alias="ts")
):
    """Upload multiple files and save to timestamp folder. Files are not processed immediately."""
    try:
        client = getattr(request, "client", None)
        client_str = f"{getattr(client, 'host', '?')}:{getattr(client, 'port', '?')}" if client else "?"
        print(f"[UPLOAD] from {client_str} ct={request.headers.get('content-type')} len={request.headers.get('content-length')}")
    except Exception:
        pass
    
    # Create or reuse timestamp folder
    def is_valid_ts(ts: str) -> bool:
        return len(ts) == 15 and ts[8] == '_' and ts.replace('_','').isdigit()
    print(f"[UPLOAD] form timestamp={ts_override} query ts={ts_q}")
    if ts_override and is_valid_ts(ts_override):
        ts = ts_override
    elif ts_q and is_valid_ts(ts_q):
        ts = ts_q
    else:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    upload_folder = UPLOAD_DIR / ts
    os.makedirs(upload_folder, exist_ok=True)
    
    saved_files = []
    total_size = 0
    
    for file in files:
        ext = Path(file.filename).suffix.lower()
        file_id = uuid.uuid4().hex[:8]
        # Keep original filename for better identification
        safe_filename = f"{file_id}_{file.filename}"
        dest = upload_folder / safe_filename
        
        # Save file with progress tracking for large files
        try:
            file_size = 0
            with open(dest, "wb") as f:
                while chunk := await file.read(1024 * 1024):  # Read 1MB at a time
                    f.write(chunk)
                    file_size += len(chunk)
            
            saved_files.append({
                "original_name": file.filename,
                "saved_name": safe_filename,
                "size": file_size,
                "type": ext
            })
            total_size += file_size
            print(f"[INFO] Saved {file.filename} ({file_size / 1024 / 1024:.2f} MB)")
        except Exception as e:
            print(f"[ERROR] Failed to save {file.filename}: {e}")
            continue
    
    # Save metadata
    metadata = {
        "timestamp": ts,
        "category": category,
        "files": saved_files,
        "total_size": total_size,
        "upload_time": datetime.now().isoformat(),
        "processed": "pending"
    }
    
    meta_path = upload_folder / "_metadata.json"
    # 合并写：如果已存在元数据则追加文件与累加大小
    if meta_path.exists():
        try:
            with open(meta_path, "r", encoding="utf-8") as f:
                old = json.load(f)
            old_files = old.get("files", [])
            old_total = old.get("total_size", 0)
            old_category = old.get("category", category)
            merged = {
                "timestamp": old.get("timestamp", ts),
                "category": category or old_category,
                "files": old_files + saved_files,
                "total_size": old_total + total_size,
                "upload_time": old.get("upload_time", datetime.now().isoformat()),
                "processed": old.get("processed", "pending")
            }
            metadata = merged
        except Exception:
            pass
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(metadata, f, ensure_ascii=False, indent=2)
    
    return JSONResponse({
        "success": True,
        "timestamp": ts,
        "folder": str(upload_folder),
        "files_count": len(saved_files),
        "total_size_mb": round(total_size / 1024 / 1024, 2),
        "message": f"成功上传 {len(saved_files)} 个文件到文件夹 {ts}"
    })


@app.post("/upload/{ts}")
async def upload_with_ts(
    ts: str,
    request: Request,
    files: List[UploadFile] = File(...),
    category: str = Form("未分类")
):
    """Upload files into a specific timestamp folder provided in path."""
    try:
        client = getattr(request, "client", None)
        client_str = f"{getattr(client, 'host', '?')}:{getattr(client, 'port', '?')}" if client else "?"
        print(f"[UPLOAD-PATH] from {client_str} ct={request.headers.get('content-type')} len={request.headers.get('content-length')} ts={ts}")
    except Exception:
        pass

    def is_valid_ts(v: str) -> bool:
        return isinstance(v, str) and len(v) == 15 and v[8] == '_' and v.replace('_','').isdigit()

    ts_use = ts if is_valid_ts(ts) else datetime.now().strftime("%Y%m%d_%H%M%S")
    upload_folder = UPLOAD_DIR / ts_use
    os.makedirs(upload_folder, exist_ok=True)

    saved_files = []
    total_size = 0
    for file in files:
        ext = Path(file.filename).suffix.lower()
        file_id = uuid.uuid4().hex[:8]
        safe_filename = f"{file_id}_{file.filename}"
        dest = upload_folder / safe_filename
        try:
            file_size = 0
            with open(dest, "wb") as f:
                while chunk := await file.read(1024 * 1024):
                    f.write(chunk)
                    file_size += len(chunk)
            saved_files.append({
                "original_name": file.filename,
                "saved_name": safe_filename,
                "size": file_size,
                "type": ext
            })
            total_size += file_size
        except Exception as e:
            print(f"[ERROR] Failed to save {file.filename}: {e}")
            continue

    meta_path = upload_folder / "_metadata.json"
    metadata = {
        "timestamp": ts_use,
        "category": category,
        "files": saved_files,
        "total_size": total_size,
        "upload_time": datetime.now().isoformat(),
        "processed": "pending"
    }
    if meta_path.exists():
        try:
            with open(meta_path, "r", encoding="utf-8") as f:
                old = json.load(f)
            metadata["files"] = old.get("files", []) + saved_files
            metadata["total_size"] = old.get("total_size", 0) + total_size
            metadata["category"] = category or old.get("category", category)
            metadata["upload_time"] = old.get("upload_time", metadata["upload_time"])
            metadata["processed"] = old.get("processed", "pending")
        except Exception:
            pass
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(metadata, f, ensure_ascii=False, indent=2)

    return JSONResponse({
        "success": True,
        "timestamp": ts_use,
        "folder": str(upload_folder),
        "files_count": len(saved_files),
        "total_size_mb": round(total_size / 1024 / 1024, 2),
        "message": f"成功上传 {len(saved_files)} 个文件到文件夹 {ts_use}"
    })


@app.get("/api/upload-folders")
async def list_upload_folders():
    """List all upload folders. If metadata missing, infer from files."""
    folders = []

    def summarize_folder(folder: Path) -> dict:
        meta_path = folder / "_metadata.json"
        if meta_path.exists():
            try:
                with open(meta_path, "r", encoding="utf-8") as f:
                    metadata = json.load(f)
                
                # Compatibility for old `processed: True/False`
                processed_status = metadata.get("processed")
                if isinstance(processed_status, bool):
                    processed_status = "completed" if processed_status else "pending"

                return {
                    "timestamp": folder.name,
                    "category": metadata.get("category", "未分类"),
                    "files_count": len(metadata.get("files", [])),
                    "total_size_mb": round(metadata.get("total_size", 0) / 1024 / 1024, 2),
                    "upload_time": metadata.get("upload_time", ""),
                    "processed": processed_status or "pending",
                    "files": metadata.get("files", []),
                    "blog_slug": metadata.get("blog_slug"),
                    "error": metadata.get("error")
                }
            except Exception as e:
                print(f"[ERROR] Failed to read metadata for {folder.name}: {e}")
        # No metadata: infer by scanning files
        file_entries = []
        total_size = 0
        for p in folder.iterdir():
            if p.is_file() and p.name != "_metadata.json":
                size = p.stat().st_size
                total_size += size
                file_entries.append({
                    "original_name": p.name,
                    "saved_name": p.name,
                    "size": size,
                    "type": p.suffix.lower()
                })
        try:
            mtime = datetime.fromtimestamp(folder.stat().st_mtime).isoformat()
        except Exception:
            mtime = ""
        return {
            "timestamp": folder.name,
            "category": "未分类",
            "files_count": len(file_entries),
            "total_size_mb": round(total_size / 1024 / 1024, 2),
            "upload_time": mtime,
            "processed": "pending",
            "files": file_entries
        }

    for folder in sorted(UPLOAD_DIR.iterdir(), reverse=True):
        try:
            if folder.is_dir():
                print(f"[LIST] folder: {folder.name}")
                info = summarize_folder(folder)
                print(f"[LIST] added: {info.get('timestamp')} files={info.get('files_count')}")
                folders.append(info)
        except Exception as e:
            print(f"[LIST] error on {folder}: {e}")

    return JSONResponse({"folders": folders})


def _update_meta_status(timestamp: str, status: str, error_msg: Optional[str] = None, slug: Optional[str] = None):
    """Helper to update the status in the metadata file."""
    meta_path = UPLOAD_DIR / timestamp / "_metadata.json"
    if not meta_path.exists():
        return

    try:
        with open(meta_path, "r+", encoding="utf-8") as f:
            data = json.load(f)
            data["processed"] = status
            if error_msg:
                data["error"] = error_msg
            if slug:
                data["blog_slug"] = slug
            data["processed_time"] = datetime.now().isoformat()
            
            f.seek(0)
            f.truncate()
            json.dump(data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"[ERROR] Failed to update metadata status for {timestamp}: {e}")


def _generate_blog_task(timestamp: str, category: Optional[str] = None):
    """Generate blog post from a specific upload folder (long-running task)"""
    
    upload_folder = UPLOAD_DIR / timestamp
    if not upload_folder.exists():
        print(f"[GEN-TASK] Folder not found: {upload_folder}")
        _update_meta_status(timestamp, "failed", "文件夹不存在")
        return
    
    try:
        # Read or build metadata
        meta_path = upload_folder / "_metadata.json"
        metadata = None
        if meta_path.exists():
            try:
                with open(meta_path, "r", encoding="utf-8") as f:
                    metadata = json.load(f)
            except Exception as e:
                print(f"[WARN] Failed to parse metadata for {timestamp}: {e}")
                metadata = None
        if metadata is None:
            file_entries = []
            total_size = 0
            for p in upload_folder.iterdir():
                if p.is_file() and p.name != "_metadata.json":
                    size = p.stat().st_size
                    total_size += size
                    file_entries.append({
                        "original_name": p.name, "saved_name": p.name, "size": size, "type": p.suffix.lower()
                    })
            metadata = {
                "timestamp": timestamp, "category": category or "未分类", "files": file_entries,
                "total_size": total_size, "upload_time": datetime.now().isoformat(), "processed": "running"
            }
            with open(meta_path, "w", encoding="utf-8") as f:
                json.dump(metadata, f, ensure_ascii=False, indent=2)
        
        final_category = category if category else metadata.get("category", "未分类")
        
        # Classify files
        paper_files, ppt_files, note_files = [], [], []
        for file_info in metadata.get("files", []):
            file_path = upload_folder / file_info["saved_name"]
            if not file_path.exists(): continue
            
            fname_lower = file_info["original_name"].lower()
            ext = file_info["type"]
            
            if ext == ".pptx" or "ppt" in fname_lower or "slide" in fname_lower:
                ppt_files.append((file_info["original_name"], file_path))
            elif "讨论" in fname_lower or "纪要" in fname_lower or "note" in fname_lower or "discussion" in fname_lower:
                note_files.append((file_info["original_name"], file_path))
            elif ext in [".pdf", ".docx", ".doc"]:
                if any(kw in fname_lower for kw in ["slide", "ppt", "presentation"]):
                     ppt_files.append((file_info["original_name"], file_path))
                else:
                     paper_files.append((file_info["original_name"], file_path))
            else:
                note_files.append((file_info["original_name"], file_path))
        
        print(f"[INFO] Task {timestamp}: {len(paper_files)} papers, {len(ppt_files)} ppts, {len(note_files)} notes")
        
        slug = uuid.uuid4().hex[:10]
        date = datetime.now().strftime("%Y-%m-%d")
        
        paper_text = ""
        if paper_files:
            texts = [f"### 文献: {fname}\n{parse_file(path)[:5000]}" for fname, path in paper_files if parse_file(path)]
            paper_text = "\n\n".join(texts)

        ppt_text, ppt_images = "", []
        if ppt_files:
            texts = [f"### PPT: {fname}\n{parse_file(path)[:3000]}" for fname, path in ppt_files if parse_file(path)]
            ppt_text = "\n\n".join(texts)
            for fname, path in ppt_files:
                if path.suffix.lower() == ".pptx":
                    ppt_images.extend(pptx_to_images(path, f"{slug}_ppt"))
                elif path.suffix.lower() == ".pdf":
                    ppt_images.extend(pdf_to_images(path, f"{slug}_ppt_pdf"))

        notes_text, note_images = "", []
        if note_files:
            texts = [f"### 讨论: {fname}\n{parse_file(path)[:3000]}" for fname, path in note_files if parse_file(path)]
            notes_text = "\n\n".join(texts)
            for fname, path in note_files:
                if path.suffix.lower() == ".pdf":
                    note_images.extend(pdf_to_images(path, f"{slug}_notes"))

        structured_content = f"""
请基于以下内容生成一篇结构化的学术博客文章:
## 一、文献原文摘要
{paper_text[:4000] if paper_text else "无"}
## 二、PPT分享内容
{ppt_text[:3000] if ppt_text else "无"}
## 三、讨论纪要
{notes_text[:3000] if notes_text else "无"}
请为每个部分生成清晰、独立的摘要。请按以下格式生成完整的博客文章:
1.  **标题**: [一个简洁、吸引人的标题]
2.  **摘要**: [对整篇文章的2-3句话总结]
3.  **标签**: [3-5个相关的关键词，用逗号分隔]
4.  **分类**: [{final_category}]
5.  **正文**:
    ## 📄 文献原文概述
    [这里是文献的核心内容概述]
    ## 🎯 PPT分享内容
    [这里是PPT的重点内容解析]
    ## 💬 讨论纪要
    [这里是讨论的要点总结]
"""
        
        analysis = analyzer.call_deepseek_blog(structured_content, category=final_category)
        
        content_parts = []
        full_content = analysis.get("content", "")

        def get_section(start_marker, end_marker=None):
            try:
                content = full_content.split(start_marker)[1]
                if end_marker:
                    content = content.split(end_marker)[0]
                return content.strip()
            except IndexError:
                return ""

        if paper_text:
            paper_summary = get_section("## 📄 文献原文概述", "## 🎯 PPT分享内容") or get_section("## 一、", "## 二、")
            content_parts.append(f"## 📄 文献原文概述\n\n{paper_summary or 'AI未能生成文献概述。'}")

        if ppt_files:
            ppt_summary = get_section("## 🎯 PPT分享内容", "## 💬 讨论纪要") or get_section("## 二、", "## 三、")
            ppt_section = f"## 🎯 PPT分享内容\n\n{ppt_summary}\n\n"
            if ppt_images:
                ppt_section += '<div class="image-gallery ppt-gallery">\n'
                for i, img_path in enumerate(ppt_images):
                    ppt_section += f'<figure><img src="{img_path}" alt="PPT幻灯片 {i+1}" loading="lazy"><figcaption>图 {i+1}: PPT 幻灯片</figcaption></figure>\n'
                ppt_section += '</div>'
            content_parts.append(ppt_section)
        
        if notes_text:
            notes_summary = get_section("## 💬 讨论纪要") or get_section("## 三、")
            notes_html = markdown.markdown(notes_summary or notes_text[:1500], extensions=['nl2br', 'fenced_code'])
            discussion_section = f'## 💬 讨论纪要\n\n<div class="discussion-notes">\n{notes_html}\n</div>'
            if note_images:
                discussion_section += '\n<div class="image-gallery notes-gallery">\n'
                for i, img_path in enumerate(note_images):
                    discussion_section += f'<figure><img src="{img_path}" alt="讨论插图 {i+1}" loading="lazy"><figcaption>图 {i+1}: 讨论插图</figcaption></figure>\n'
                discussion_section += '</div>'
            content_parts.append(discussion_section)
        
        content_html = "\n\n".join(content_parts) if content_parts else markdown.markdown(analysis.get("content", "无内容"))
        
        tpl = env.get_template("post.html")
        html = tpl.render(
            title=analysis.get("title", "未命名"), summary=analysis.get("summary", ""), tags=analysis.get("tags", []),
            content=content_html, category=analysis.get("category", final_category), date=date
        )
        (SITE_DIR / f"{slug}.html").write_text(html, encoding="utf-8")
        
        post_meta = {
            "title": analysis.get("title", "未命名"), "date": date, "category": analysis.get("category", final_category),
            "tags": analysis.get("tags", []), "summary": analysis.get("summary", ""),
            "has_paper": bool(paper_files), "has_ppt": bool(ppt_files), "has_notes": bool(note_files),
            "paper_images": [], "ppt_images": ppt_images, "note_images": note_images, "source_folder": timestamp
        }
        with open(SITE_DIR / f"{slug}.json", "w", encoding="utf-8") as mf:
            json.dump(post_meta, mf, ensure_ascii=False, indent=2)
        
        _update_meta_status(timestamp, "completed", slug=slug)
        print(f"[SUCCESS] Task {timestamp} completed. Blog: /site/{slug}.html")

    except Exception as e:
        print(f"[ERROR] Blog generation task for {timestamp} failed: {e}")
        traceback.print_exc()
        _update_meta_status(timestamp, "failed", error_msg=str(e))


@app.post("/api/generate-blog/{timestamp}")
async def generate_blog_from_folder(timestamp: str, background_tasks: BackgroundTasks, category: Optional[str] = Query(None)):
    """Generate blog post from a specific upload folder in the background"""
    
    upload_folder = UPLOAD_DIR / timestamp
    if not upload_folder.exists():
        return JSONResponse({"error": "文件夹不存在"}, status_code=404)

    _update_meta_status(timestamp, "running")
    background_tasks.add_task(_generate_blog_task, timestamp, category)
    
    return JSONResponse({"message": "博文生成任务已在后台开始", "timestamp": timestamp}, status_code=202)


@app.get("/api/generation-status/{timestamp}")
async def get_generation_status(timestamp: str):
    """Check the status of a blog generation task."""
    meta_path = UPLOAD_DIR / timestamp / "_metadata.json"
    if not meta_path.exists():
        return JSONResponse({"error": "任务不存在"}, status_code=404)
    
    try:
        with open(meta_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        
        status = data.get("processed", "pending")
        if isinstance(status, bool): # backward compatibility
            status = "completed" if status else "pending"
            
        return JSONResponse({
            "timestamp": timestamp,
            "status": status,
            "error": data.get("error"),
            "blog_slug": data.get("blog_slug"),
            "url": f"/site/{data.get('blog_slug')}.html" if data.get('blog_slug') else None
        })
    except Exception as e:
        return JSONResponse({"status": "unknown", "error": f"无法读取状态: {e}"}, status_code=500)


@app.get("/health")
async def health():
    return {"status": "ok"}


@app.get("/admin", response_class=HTMLResponse)
async def admin(request: Request):
    posts = []
    for f in sorted(SITE_DIR.glob("*.json"), key=lambda x: x.stat().st_mtime, reverse=True):
        try:
            with open(f, "r", encoding="utf-8") as mf:
                data = json.load(mf)
            posts.append({
                "name": data.get("title") or f.stem,
                "title": data.get("title") or f.stem,
                "slug": data.get("slug") or f.stem,
                "category": data.get("category", "未分类"),
                "date": data.get("date", "未知")
            })
        except Exception:
            posts.append({
                "name": f.stem,
                "title": f.stem,
                "slug": f.stem,
                "category": "未分类",
                "date": "未知"
            })
    tpl = env.get_template("admin.html")
    openwebui = os.getenv("OPENWEBUI_URL", "http://localhost:3000")
    return tpl.render(posts=posts, openwebui=openwebui)


@app.get("/post/{slug}")
async def get_post(slug: str):
    meta_path = SITE_DIR / f"{slug}.json"
    if not meta_path.exists():
        return JSONResponse({"error": "not found"}, status_code=404)
    with open(meta_path, "r", encoding="utf-8") as mf:
        data = json.load(mf)
    return JSONResponse(data)


@app.post("/review/{slug}")
async def review_post(slug: str, request: Request):
    meta_path = SITE_DIR / f"{slug}.json"
    if not meta_path.exists():
        return JSONResponse({"error": "not found"}, status_code=404)
    payload = await request.json()
    title = payload.get("title")
    summary = payload.get("summary")
    content = payload.get("content")
    tags = payload.get("tags") or []
    # update meta json
    try:
        # Read existing metadata to preserve date and category
        with open(meta_path, "r", encoding="utf-8") as mf:
            old_data = json.load(mf)
        
        # Update with new values
        with open(meta_path, "w", encoding="utf-8") as mf:
            json.dump({
                "title": title,
                "summary": summary,
                "tags": tags,
                "content": content,
                "slug": slug,
                "date": old_data.get("date", datetime.now().strftime("%Y-%m-%d")),
                "category": old_data.get("category", "未分类")
            }, mf, ensure_ascii=False, indent=2)
        
        # re-render html with markdown conversion
        content_html = markdown.markdown(
            content,
            extensions=['fenced_code', 'tables', 'nl2br']
        )
        tpl = env.get_template("post.html")
        html = tpl.render(
            title=title,
            summary=summary,
            tags=tags,
            content=content_html,
            date=old_data.get("date", datetime.now().strftime("%Y-%m-%d")),
            category=old_data.get("category", "未分类")
        )
        out_path = SITE_DIR / f"{slug}.html"
        out_path.write_text(html, encoding="utf-8")
        return JSONResponse({"ok": True, "url": f"/site/{slug}.html"})
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.post("/import")
async def import_file(request: Request):
    """Import a file already present on the server under the `uploads/` directory.
    POST JSON {"filename": "your.pdf"}
    This helps avoid HTTP upload size limits imposed by some proxies.
    """
    payload = await request.json()
    filename = payload.get("filename")
    if not filename:
        return JSONResponse({"error": "filename required"}, status_code=400)

    # Sanitize path: only allow files inside UPLOAD_DIR
    candidate = UPLOAD_DIR / Path(filename).name
    if not candidate.exists():
        return JSONResponse({"error": "file not found on server"}, status_code=404)

    # parse
    text = parse_file(candidate)
    analysis = analyzer.call_deepseek(text)

    slug = analysis.get("slug") or Path(filename).stem
    
    # Convert markdown to HTML
    content_html = markdown.markdown(
        analysis.get("content", ""),
        extensions=['fenced_code', 'tables', 'nl2br']
    )
    
    tpl = env.get_template("post.html")
    html = tpl.render(
        title=analysis.get("title", filename),
        summary=analysis.get("summary", ""),
        tags=analysis.get("tags", []),
        content=content_html
    )
    out_path = SITE_DIR / f"{slug}.html"
    out_path.write_text(html, encoding="utf-8")

    try:
        meta_path = SITE_DIR / f"{slug}.json"
        with open(meta_path, "w", encoding="utf-8") as mf:
            json.dump({"title": analysis.get("title", filename), "summary": analysis.get("summary",""), "tags": analysis.get("tags",[]), "content": analysis.get("content",""), "slug": slug}, mf, ensure_ascii=False, indent=2)
    except Exception:
        pass

    return JSONResponse({"ok": True, "url": f"/site/{slug}.html", "slug": slug})


@app.delete("/post/{slug}")
async def delete_post(slug: str):
    """Delete a post (both HTML and JSON files)."""
    html_path = SITE_DIR / f"{slug}.html"
    json_path = SITE_DIR / f"{slug}.json"
    
    deleted = []
    if html_path.exists():
        html_path.unlink()
        deleted.append("html")
    if json_path.exists():
        json_path.unlink()
        deleted.append("json")
    
    if deleted:
        return JSONResponse({"ok": True, "deleted": deleted})
    else:
        return JSONResponse({"error": "not found"}, status_code=404)


@app.post("/api/ai-rewrite")
async def ai_rewrite(request: Request):
    """Use AI to improve/rewrite existing blog content."""
    payload = await request.json()
    title = payload.get("title", "")
    summary = payload.get("summary", "")
    content = payload.get("content", "")
    tags = payload.get("tags", "")
    
    # Combine current content and ask AI to improve it
    combined = f"标题: {title}\n\n摘要: {summary}\n\n正文:\n{content}\n\n标签: {tags}"
    
    try:
        # Use the blog generator with rewrite instruction
        result = analyzer.call_deepseek_blog(combined, category="优化版")
        return JSONResponse({
            "title": result.get("title") or title,
            "summary": result.get("summary") or summary,
            "content": result.get("content") or content,
            "tags": result.get("tags") or []
        })
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


    return JSONResponse({"ok": True, "url": f"/site/{slug}.html"})
